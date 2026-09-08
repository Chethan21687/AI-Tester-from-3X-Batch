#Test Analyst
#A senior QA with 15 years (JIRA MD) of experience in testing web applications, mobile applications, and APIs. It will create analyze the requirements and suggest exactly 10 testcases 


from crewai import Agent,Task,Crew
from crewai import LLM
from dotenv import load_dotenv
import os

# OpenAI - GROQ API Key


# Step 0 - Set up the Brain
# Step 1. - Define the Agent (identity)
# Step 2. - Give the Task to the Agent
# Step 3. Add them to the Crew
# Step 4. Kick Off Agent.

# Prompt vs Skill vs AI Agent

# We need use the GROQ gpt-oss-120b model


def build_groq_llm():
    # Step 0 - Set up the Brain (Groq LLM)
    load_dotenv(override=True)  # reads the .env file in this folder (override stale env vars like GROQ_MODEL)

    # Groq exposes an OpenAI-compatible API, so we use the "openai/" provider
    # prefix with a custom base_url pointing at Groq.
    # GROQ_MODEL in .env = openai/gpt-oss-120b (exact ID from groq.com console)
    return LLM(
        model=f"openai/{os.getenv('GROQ_MODEL')}",
        base_url=os.getenv('GROQ_BASE_URL'),
        api_key=os.getenv('GROQ_API_KEY'),
        max_tokens=4096,
    )


def build_crew(requirements: str = "", verbose: bool = True, count: int = 10) -> Crew:
    """Create the QA Crew. Pass requirements text to scope the generated test cases."""
    # Step 1 - Define the Agent (identity)
    qaagent = Agent(
        role="Test Analyst",
        goal=f"Create {count} test cases for a web application based on the provided requirements.",
        backstory=(f"A senior QA with 15 years of experience in testing web applications, mobile applications, and APIs. "
                    f"It will analyze the requirements and suggest exactly {count} test cases."),
        llm=build_groq_llm(),
        verbose=verbose
    )

    # Step 2 - Give the Task to the Agent
    description = f"Create {count} test cases for a web application based on the provided requirements."
    if requirements and requirements.strip():
        description = (
            f"Create {count} test cases for a web application based on the provided requirements.\n\n"
            f"Requirements:\n{requirements.strip()}"
        )
    task = Task(
        description=description,
        expected_output=f"A list of {count} test cases with clear descriptions, expected results, and any necessary preconditions.",
        agent=qaagent,
        verbose=verbose
    )

    # Step 3 - Add them to the Crew
    return Crew(
        name="QA Crew",
        agents=[qaagent],
        tasks=[task],
        verbose=verbose
    )


def build_docx(markdown_text):
    """Convert the crew's markdown output (incl. the test-case table) into a Word Document."""
    import re
    from docx import Document

    def clean(s):
        s = s.replace('<br>', '\n')
        s = re.sub(r'\*\*(.+?)\*\*', r'\1', s)   # bold
        s = re.sub(r'\*(.+?)\*', r'\1', s)       # italic
        s = s.replace('`', '')
        return s.strip()

    doc = Document()
    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # --- markdown table (consecutive '|' rows) ---
        if stripped.startswith('|') and stripped.endswith('|') and '|' in stripped[1:]:
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                rows.append(cells)
                i += 1
            rows = [r for r in rows if not all(re.fullmatch(r':?-{2,}:?', c.replace(' ', '')) for c in r)]
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    for ci in range(ncols):
                        cell_text = clean(r[ci]) if ci < len(r) else ''
                        cell = table.rows[ri].cells[ci]
                        first = True
                        for para_text in cell_text.split('\n'):
                            if not para_text:
                                continue
                            p = cell.paragraphs[0] if first else cell.add_paragraph()
                            first = False
                            run = p.add_run(para_text)
                            if ri == 0:
                                run.bold = True
                doc.add_paragraph()
            continue

        # --- headings / paragraphs / bullets ---
        if re.match(r'^#{1,6}\s', stripped):
            level = len(stripped) - len(stripped.lstrip('#'))
            doc.add_heading(clean(stripped.lstrip('#')), level=min(level, 4))
        elif stripped.startswith(('- ', '* ')):
            doc.add_paragraph(clean(stripped[2:]), style='List Bullet')
        elif stripped:
            doc.add_paragraph(clean(stripped))
        i += 1

    return doc


def save_result_to_docx(markdown_text, path):
    """Save the crew's output as a Word .docx file."""
    doc = build_docx(markdown_text)
    doc.save(path)
    print(f"\nSaved output to: {path}")


# Step 4 - Kick Off Agent.
if __name__ == "__main__":
    crew = build_crew()
    result = crew.kickoff()
    print(result)
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "TestForge_Test_Cases.docx")
    save_result_to_docx(str(result), output_path)
