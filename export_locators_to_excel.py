"""
Convert a locator markdown table (| Element | Locator | Strategy | Needs XPath? |)
into an .xlsx workbook, with XPath-fallback rows highlighted.

Usage: python export_locators_to_excel.py <input.md> <output.xlsx>
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
XPATH_FILL = PatternFill("solid", fgColor="FCE4D6")
HEADER_FONT = Font(bold=True)


def _split_row(raw, ncols):
    """Split a (possibly multi-line) '| a | b | c |' chunk into ncols cells, or None if incomplete."""
    cells = [c.strip().strip("`") for c in raw.strip().strip("|").split("|")]
    return cells if len(cells) == ncols else None


def parse_table(path):
    """Parse a markdown table where cell values may contain literal newlines.
    A new row starts on a line beginning with '|' once the previous buffered
    row already split cleanly into the expected column count; otherwise the
    line is a continuation of the previous cell's text."""
    with open(path, encoding="utf-8") as fh:
        lines = [l.rstrip("\n") for l in fh]

    header_line = next(l for l in lines if l.strip().startswith("|"))
    header = [c.strip().strip("`") for c in header_line.strip().strip("|").split("|")]
    ncols = len(header)

    rows, buf = [], ""
    started = False
    for line in lines:
        if not started:
            if line is header_line:
                started = True
            continue
        if set(line.replace("|", "").strip()) <= {"-"} and "-" in line:
            continue  # separator row (|---|---|)
        if line.strip().startswith("|") and (not buf or _split_row(buf, ncols)):
            if buf:
                rows.append(_split_row(buf, ncols))
            buf = line
        else:
            buf = f"{buf}\n{line}" if buf else line
    if buf:
        complete = _split_row(buf, ncols)
        if complete:
            rows.append(complete)
    return header, rows


def main():
    if len(sys.argv) != 3:
        print("Usage: python export_locators_to_excel.py <input.md> <output.xlsx>")
        sys.exit(1)
    in_path, out_path = sys.argv[1], sys.argv[2]

    header, data_rows = parse_table(in_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Locators"
    ws.append(header)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    needs_xpath_col = header.index("Needs XPath?")
    for row in data_rows:
        ws.append(row)
        if row[needs_xpath_col].strip().lower() == "yes":
            for cell in ws[ws.max_row]:
                cell.fill = XPATH_FILL

    widths = [40, 70, 18, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    wb.save(out_path)
    print(f"Wrote {len(data_rows)} rows to {out_path} ({sum(1 for r in data_rows if r[needs_xpath_col].lower()=='yes')} flagged as XPath fallback)")


if __name__ == "__main__":
    main()
