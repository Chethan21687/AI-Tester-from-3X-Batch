"""
Re-run the focused amazon.in locator extraction and write straight to .xlsx —
skips the markdown round-trip because some accessible names contain literal
'|' / newline characters that collide with markdown table syntax.

Usage: python export_amazon_xlsx.py <output.xlsx>
"""
import sys
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from amazon_main_locators import extract

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
XPATH_FILL = PatternFill("solid", fgColor="FCE4D6")
HEADER_FONT = Font(bold=True)
HEADER = ["Element", "Locator", "Strategy", "Needs XPath?"]


def main():
    out_path = sys.argv[1] if len(sys.argv) > 1 else "amazon_locators.xlsx"
    rows = extract()

    wb = Workbook()
    ws = wb.active
    ws.title = "Locators"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for element, locator, strategy, needs_xpath in rows:
        ws.append([element.replace("\n", " "), locator, strategy, needs_xpath])
        if needs_xpath == "Yes":
            for cell in ws[ws.max_row]:
                cell.fill = XPATH_FILL

    for col, width in zip("ABCD", (45, 75, 18, 14)):
        ws.column_dimensions[col].width = width

    wb.save(out_path)
    flagged = sum(1 for r in rows if r[3] == "Yes")
    print(f"Wrote {len(rows)} rows to {out_path} ({flagged} flagged as XPath fallback)")


if __name__ == "__main__":
    main()
